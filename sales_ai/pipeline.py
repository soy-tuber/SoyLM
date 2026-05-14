"""Ingestion: raw file -> extracted text -> LLM-structured JSON -> child row.

Each ingestion is bound to a `deal_id` (an existing 案件マスタ row).
"""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from typing import Any

from . import db_manager, llm_router, schemas


# --- text extraction ---------------------------------------------------------

def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        return _extract_pdf(data)
    if name.endswith(".docx"):
        return _extract_docx(data)
    if name.endswith(".pptx"):
        return _extract_pptx(data)
    if name.endswith(".xlsx"):
        return _extract_xlsx(data)
    if name.endswith(".csv"):
        return data.decode("utf-8", errors="replace")
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    return "\n\n".join(page.get_text() for page in doc)


def _extract_docx(data: bytes) -> str:
    import docx
    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text:
                        out.append(text)
    return "\n".join(out)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []
    for sheet in wb.sheetnames:
        out.append(f"--- Sheet: {sheet} ---")
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                out.append("\t".join(cells))
    return "\n".join(out)


# --- LLM structuring ---------------------------------------------------------

@dataclass
class StructureResult:
    record: dict[str, Any]
    raw_response: str


def _build_structuring_prompt(schema: schemas.ChildSchema) -> str:
    fields_desc = "\n".join(f"- {c} ({t})" for c, t in schema.columns)
    keys_csv = ", ".join(c for c, _ in schema.columns)
    return f"""あなたはデータ構造化アシスタントです。
入力テキストから、指定スキーマに従って JSON オブジェクトを 1 つだけ出力します。

## スキーマ: {schema.name} ({schema.display})
{schema.prompt_hint}

## フィールド (全て出力。値が不明な場合は空文字列 "" または 0)
{fields_desc}

## 出力ルール
- JSON オブジェクト 1 個のみ出力。前後の説明文・コードフェンスは禁止。
- キーは厳密に次のみ: {keys_csv}
- 値は日本語の自然文で簡潔に。リスト・辞書は使わず文字列にまとめる。
- INTEGER 型は数値 (不明なら 0)。
- 長文フィールド (full_text, script など) は入力本文を可能な限り保持する。
"""


_JSON_RE = re.compile(r"\{.*\}", re.S)


def _parse_json(text: str) -> dict[str, Any]:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.S)
    try:
        return json.loads(text)
    except Exception:
        m = _JSON_RE.search(text)
        if not m:
            raise ValueError("LLM did not return JSON")
        return json.loads(m.group(0))


def structure_text(
    schema_name: str,
    text: str,
    provider: str,
    model: str,
    api_key: str,
    max_chars: int = 30000,
) -> StructureResult:
    sc = schemas.get_child(schema_name)
    truncated = text[:max_chars]
    system = _build_structuring_prompt(sc)
    response = llm_router.chat(
        provider=provider, model=model, api_key=api_key,
        system=system,
        messages=[{"role": "user", "content": truncated}],
        max_tokens=4096,
    )
    record = _parse_json(response)
    for col, sql_type in sc.columns:
        if sql_type == "INTEGER" and col in record:
            try:
                record[col] = int(re.sub(r"[^\d-]", "", str(record[col])) or 0)
            except ValueError:
                record[col] = 0
    return StructureResult(record=record, raw_response=response)


def ingest(
    schema_name: str,
    deal_id: int,
    filename: str,
    data: bytes,
    provider: str,
    model: str,
    api_key: str,
) -> dict[str, Any]:
    text = extract_text(filename, data)
    if not text.strip():
        raise ValueError("ファイルからテキストを抽出できませんでした")
    result = structure_text(schema_name, text, provider, model, api_key)
    row_id = db_manager.insert_child(schema_name, deal_id, result.record, source_file=filename)
    return {"id": row_id, "deal_id": deal_id, "record": result.record, "extracted_chars": len(text)}
